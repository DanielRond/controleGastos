"""Testes da geração de planilhas (generator.py)."""

from pathlib import Path

from openpyxl import load_workbook

from controlegastos.generator import gerar_planilha
from controlegastos.loader import carregar_json

RAIZ = Path(__file__).resolve().parent.parent
EXEMPLO = RAIZ / "data" / "entradas" / "exemplo.json"


def test_gera_planilha_com_abas_esperadas(tmp_path: Path) -> None:
    livros = carregar_json(EXEMPLO)
    destino = gerar_planilha(livros, tmp_path / "planilha.xlsx")

    assert destino.exists()
    workbook = load_workbook(destino)
    assert workbook.sheetnames == ["Resumo Geral", "2025-01", "2025-02", "Gráficos"]


def test_resumo_geral_tem_cabecalho(tmp_path: Path) -> None:
    livros = carregar_json(EXEMPLO)
    destino = gerar_planilha(livros, tmp_path / "planilha.xlsx")

    workbook = load_workbook(destino)
    resumo = workbook["Resumo Geral"]
    assert resumo["A1"].value == "Ano"
    assert resumo["B1"].value == "Mês"


def test_aba_do_mes_tem_linha_saldo_livre(tmp_path: Path) -> None:
    livros = carregar_json(EXEMPLO)
    destino = gerar_planilha(livros, tmp_path / "planilha.xlsx")

    workbook = load_workbook(destino)
    planilha = workbook["2025-01"]
    assert planilha["A1"].value == "Data"

    saldo_celula = None
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if linha[2] == "Saldo Livre":
            saldo_celula = linha[3]
    assert saldo_celula is not None
    assert round(float(saldo_celula), 2) == 2470.00


def test_aba_graficos_tem_graficos(tmp_path: Path) -> None:
    livros = carregar_json(EXEMPLO)
    destino = gerar_planilha(livros, tmp_path / "planilha.xlsx")

    workbook = load_workbook(destino)
    graficos = workbook["Gráficos"]
    assert graficos._charts, "esperava ao menos um gráfico na aba Gráficos"
    assert len(graficos._charts) >= 2